"""PPTX export of an experiment's findings (FR-UI-09 stretch).

Every number on the slides comes from the ``Findings`` object (statistics code),
and every slide carries the trust indicators: fidelity mode, replications,
confidence and envelope status. Requires the optional ``python-pptx`` package.
"""

from __future__ import annotations

import io
from collections.abc import Mapping
from typing import Any

from simcore.specs.domain import MetricDef
from simcore.specs.findings import Findings

__all__ = ["findings_pptx"]


def _fmt(v: float | None, m: MetricDef | None) -> str:
    if v is None:
        return "-"
    if m is not None and m.format.kind == "percent":
        return f"{v * 100:.{m.format.decimals}f}%"
    return f"{v:,.0f}" if abs(v) >= 100 else f"{v:.3g}"


def _chart_png(findings: Findings, catalog: Mapping[str, MetricDef]) -> bytes | None:
    try:
        import vl_convert as vlc
    except ImportError:  # pragma: no cover - optional
        return None
    rows = [
        {
            "label": f"{c.variant}: "
            + (catalog[c.metric].display_name if c.metric in catalog else c.metric),
            "diff": c.paired_diff.mean,
            "lo": c.paired_diff.ci[0],
            "hi": c.paired_diff.ci[1],
            "significant": "CI excludes 0" if c.significant else "CI includes 0",
        }
        for c in findings.comparisons[:12]
    ]
    if not rows:
        return None
    spec = {
        "$schema": "https://vega.github.io/schema/vega-lite/v5.json",
        "width": 640,
        "height": 28 * len(rows),
        "data": {"values": rows},
        "encoding": {"y": {"field": "label", "type": "nominal", "title": None, "sort": None}},
        "layer": [
            {
                "mark": "rule",
                "encoding": {
                    "x": {
                        "field": "lo",
                        "type": "quantitative",
                        "title": "paired difference vs baseline",
                    },
                    "x2": {"field": "hi"},
                },
            },
            {
                "mark": {"type": "point", "filled": True, "size": 80},
                "encoding": {
                    "x": {"field": "diff", "type": "quantitative"},
                    "color": {"field": "significant", "type": "nominal"},
                },
            },
            {"mark": {"type": "rule", "strokeDash": [4, 4]}, "encoding": {"x": {"datum": 0}}},
        ],
    }
    png: bytes = vlc.vegalite_to_png(spec, scale=2)
    return png


def findings_pptx(
    findings: Findings, narrative: str, catalog: Mapping[str, MetricDef], title: str
) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    fid = findings.fidelity
    extrap = any(c.extrapolated for c in findings.comparisons)
    trust = (
        f"{fid.screen} mode{' (promoted: ' + ', '.join(fid.promoted) + ')' if fid.promoted else ''}"
        f" · n = {findings.n_replications} replications · {round(findings.confidence * 100)}% "
        f"confidence · "
        + ("EXTRAPOLATED (outside envelope)" if extrap else "within calibrated envelope")
    )

    def footer(slide: Any) -> None:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(12.3), Inches(0.4))
        box.text_frame.text = trust
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(11)

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title
    s.placeholders[1].text = f"Experiment {findings.experiment_id} vs {findings.baseline}"
    footer(s)

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Findings"
    rows = len(findings.comparisons[:14]) + 1
    table = s.shapes.add_table(
        rows, 6, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.4 * rows)
    ).table
    for j, h in enumerate(["Variant", "Metric", "Baseline", "Variant", "Δ (95% CI)", ""]):
        table.cell(0, j).text = h
    for i, c in enumerate(findings.comparisons[:14], start=1):
        m = catalog.get(c.metric)
        d = c.paired_diff
        cells = [
            c.variant,
            m.display_name if m else c.metric,
            _fmt(c.baseline.mean if c.baseline else None, m),
            _fmt(c.variant_estimate.mean if c.variant_estimate else None, m),
            f"{_fmt(d.mean, m)} ({_fmt(d.ci[0], m)} to {_fmt(d.ci[1], m)})",
            ("significant" if c.significant else "n.s.")
            + (" · extrapolated" if c.extrapolated else ""),
        ]
        for j, text in enumerate(cells):
            table.cell(i, j).text = text
            table.cell(i, j).text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    footer(s)

    png = _chart_png(findings, catalog)
    if png is not None:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "Differences vs baseline"
        s.shapes.add_picture(io.BytesIO(png), Inches(0.8), Inches(1.4), height=Inches(5.3))
        footer(s)

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Summary and caveats"
    body = s.placeholders[1].text_frame
    body.text = narrative.splitlines()[0] if narrative else ""
    for line in narrative.splitlines()[1:]:
        p = body.add_paragraph()
        p.text = line.lstrip("- ")
        p.font.size = Pt(14)
    footer(s)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
